# -*- coding: utf-8 -*-
"""
Fotos sortieren – grafische Oberfläche
Start:  Doppelklick auf die Desktop-Verknüpfung  oder  pythonw Fotosortierprogramm_v2.21.pyw
Die Versionsnummer steht im Dateinamen und in VERSION.
"""

import json
import os
import shutil
import subprocess
import time
import queue
import sys
import threading
import tkinter as tk
from pathlib import Path
from datetime import datetime
from tkinter import filedialog, messagebox, scrolledtext, simpledialog, ttk

PROGRAMMORDNER = Path(__file__).resolve().parent
sys.path.insert(0, str(PROGRAMMORDNER))
import fotosortierer_kern as kern  # noqa: E402

EINSTELLUNGSDATEI = PROGRAMMORDNER / "einstellungen.json"
ICON = PROGRAMMORDNER / "fotos_sortieren.ico"
VERSION = "2.21"

# Schriftgrössen in Punkt – hier anpassen, falls gewünscht
GROESSE = 13
SCHRIFT = ("Arial", GROESSE)
SCHRIFT_FETT = ("Arial", GROESSE, "bold")
SCHRIFT_KLEIN = ("Arial", GROESSE - 1)
SCHRIFT_ERLAEUTERUNG = ("Arial", GROESSE - 3)          # Erläuterungstexte (kleiner)
SCHRIFT_MONO = ("Consolas", GROESSE - 1)
SCHRIFT_ABSCHNITT = ("Arial", GROESSE + 4, "bold")     # Abschnittstitel 1–4 (blaue Balken)
SCHRIFT_TITEL = ("Arial", GROESSE + 13, "bold")        # Programmtitel im Kopf
SCHRIFT_UNTERTITEL = ("Arial", GROESSE - 1)
SCHRIFT_TABELLE = ("Arial", GROESSE - 3)                # Zuordnungsliste (kompakt, ohne horizontales Scrollen)
KOPF_FARBE = "#1B3F6B"                                 # Kopfbereich (dunkelblau)
BALKEN_FARBE = "#1F4E79"                               # Abschnittsbalken (Ausführen, Protokoll)
BALKEN_HELL = "#5B8DC9"                                # Balken der aufklappbaren Abschnitte 1-3
BALKEN_TEXT = "white"
ZEBRA = ("#FFFFFF", "#EEF1F5")                         # Tabellenzeilen abwechselnd
AUSWAHL = "#CFE0F3"                                    # markierte Tabellenzeile
HINWEIS_FARBE = "#555555"



class KategorienDialog(tk.Toplevel):
    """
    Thema wählen (bestehendes ergänzen oder neues anlegen) und Beschreibungen erfassen.
    kategorien: dict Name -> [Beschreibungen];  vorwahl: Name zum Bearbeiten oder None
    """
    NEU = "<< Neues Thema anlegen >>"

    def __init__(self, master, kategorien: dict, vorwahl=None, mehrsprachig=True):
        super().__init__(master)
        self.title("Thema bearbeiten" if vorwahl else "Thema ergänzen oder neu anlegen")
        self.resizable(True, True)
        self.ergebnis = None
        self.kategorien = kategorien
        self.transient(master)
        self.grab_set()

        frm = ttk.Frame(self, padding=14)
        frm.pack(fill="both", expand=True)
        frm.columnconfigure(1, weight=1)

        ttk.Label(frm, text="Thema:").grid(row=0, column=0, sticky="w", pady=(0, 4))
        self.v_thema = tk.StringVar()
        self.cb = ttk.Combobox(frm, textvariable=self.v_thema, width=40, font=SCHRIFT,
                               values=[self.NEU] + list(kategorien.keys()))
        self.cb.grid(row=0, column=1, sticky="ew", pady=(0, 4))
        self.cb.bind("<<ComboboxSelected>>", self._thema_gewechselt)

        ttk.Label(frm, text="Neuer Name:").grid(row=1, column=0, sticky="w", pady=(0, 8))
        self.name = ttk.Entry(frm, width=40, font=SCHRIFT)
        self.name.grid(row=1, column=1, sticky="ew", pady=(0, 8))

        sprache = ("auf Deutsch" if mehrsprachig else
                   "auf Deutsch (werden für das englische Modell automatisch übersetzt; englische Texte bleiben)")
        ttk.Label(frm, style="Hinweis.TLabel", justify="left",
                  text=f"Beschreibungen {sprache}, eine pro Zeile. Bestehende Zeilen bleiben erhalten,\n"
                       "neue Zeilen werden ergänzt. Mehrere Formulierungen machen die Erkennung robuster, z. B.\n"
                       "   Person mit Regenschirm\n   Menschen im Regen mit aufgespanntem Schirm").grid(
            row=2, column=0, columnspan=2, sticky="w")
        self.text = scrolledtext.ScrolledText(frm, width=70, height=10, wrap="word", font=SCHRIFT)
        self.text.grid(row=3, column=0, columnspan=2, sticky="nsew", pady=(4, 0))
        frm.rowconfigure(3, weight=1)

        kn = ttk.Frame(frm)
        kn.grid(row=4, column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(kn, text="Abbrechen", command=self.destroy).pack(side="right", padx=(6, 0))
        ttk.Button(kn, text="OK", command=self.ok).pack(side="right")
        self.bind("<Escape>", lambda e: self.destroy())

        if vorwahl and vorwahl in kategorien:
            self.v_thema.set(vorwahl)
        else:
            self.v_thema.set(self.NEU)
        self._thema_gewechselt()

    def _thema_gewechselt(self, *_):
        wahl = self.v_thema.get()
        self.text.delete("1.0", "end")
        if wahl == self.NEU:
            self.name.configure(state="normal")
            self.name.delete(0, "end")
            self.name.focus_set()
        else:
            self.name.configure(state="normal")
            self.name.delete(0, "end")
            self.name.insert(0, wahl)
            self.text.insert("1.0", "\n".join(self.kategorien.get(wahl, [])))
            self.text.focus_set()
            self.text.mark_set("insert", "end")

    def ok(self):
        wahl = self.v_thema.get()
        name = kern.sicherer_name(self.name.get().strip().replace(" ", "_"))
        zeilen = [z.strip() for z in self.text.get("1.0", "end").splitlines() if z.strip()]
        if not name or name == "Unbekannt":
            messagebox.showwarning("Fehlt", "Bitte einen Namen für das Thema eingeben.", parent=self)
            return
        if not zeilen:
            messagebox.showwarning("Fehlt", "Mindestens eine Beschreibung eingeben.", parent=self)
            return
        alt = None if wahl == self.NEU else wahl
        if alt is None and name in self.kategorien:
            if not messagebox.askyesno("Thema existiert",
                                       f"Das Thema «{name}» gibt es schon. Beschreibungen dort ergänzen?",
                                       parent=self):
                return
            zeilen = list(dict.fromkeys(self.kategorien[name] + zeilen))
        self.ergebnis = (alt, name, list(dict.fromkeys(zeilen)))
        self.destroy()


class ThemenTabelle(ttk.Frame):
    """Zweispaltige Tabelle Thema | Beschreibungen mit Zeilenumbruch und Zebra-Streifen."""

    def __init__(self, master, on_doppelklick=None, spalte1=250, **kw):
        super().__init__(master, **kw)
        self.on_doppelklick = on_doppelklick
        self.spalte1 = spalte1
        self.zeilen = []          # [(name, frame, label1, label2)]
        self.gewaehlt = None
        self.canvas = tk.Canvas(self, height=340, highlightthickness=1, highlightbackground="#B0B0B0", bg="white")
        self.sb = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=self.sb.set)
        self.canvas.grid(row=0, column=0, sticky="nsew")
        self.sb.grid(row=0, column=1, sticky="ns")
        self.rowconfigure(0, weight=1)
        self.columnconfigure(0, weight=1)
        self.inner = tk.Frame(self.canvas, bg="white")
        self.fenster = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.inner.columnconfigure(1, weight=1)
        self.canvas.bind("<Configure>", self._breite_anpassen)
        self.inner.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        for w in (self.canvas, self.inner):
            w.bind("<MouseWheel>", self._rad)
        # Kopfzeile
        kopf = tk.Frame(self.inner, bg="#D9DEE5")
        kopf.grid(row=0, column=0, columnspan=2, sticky="ew")
        kopf.columnconfigure(1, weight=1)
        tk.Label(kopf, text="Thema (Ordnername)", font=SCHRIFT_FETT, bg="#D9DEE5", anchor="w",
                 width=1, padx=8, pady=4).grid(row=0, column=0, sticky="ew", ipadx=0)
        tk.Label(kopf, text="Beschreibungen", font=SCHRIFT_FETT, bg="#D9DEE5", anchor="w",
                 padx=8, pady=4).grid(row=0, column=1, sticky="ew")
        kopf.grid_columnconfigure(0, minsize=self.spalte1)

    def _rad(self, e):
        self.canvas.yview_scroll(int(-e.delta / 120), "units")
        return "break"

    def _breite_anpassen(self, e):
        self.canvas.itemconfigure(self.fenster, width=e.width)
        wrap = max(200, e.width - self.spalte1 - 30)
        for _, _, _, l2 in self.zeilen:
            l2.configure(wraplength=wrap)

    def setzen(self, kategorien: dict):
        for _, f, _, _ in self.zeilen:
            f.destroy()
        self.zeilen = []
        self.gewaehlt = None
        wrap = max(200, self.canvas.winfo_width() - self.spalte1 - 30)
        for i, (name, b) in enumerate(kategorien.items()):
            bg = ZEBRA[i % 2]
            f = tk.Frame(self.inner, bg=bg)
            f.grid(row=i + 1, column=0, columnspan=2, sticky="ew")
            f.columnconfigure(1, weight=1)
            f.grid_columnconfigure(0, minsize=self.spalte1)
            l1 = tk.Label(f, text=name, font=SCHRIFT, bg=bg, anchor="nw", justify="left",
                          padx=8, pady=5, wraplength=self.spalte1 - 16)
            l1.grid(row=0, column=0, sticky="nsew")
            l2 = tk.Label(f, text=";  ".join(b), font=SCHRIFT, bg=bg, anchor="nw", justify="left",
                          padx=8, pady=5, wraplength=wrap)
            l2.grid(row=0, column=1, sticky="nsew")
            for w in (f, l1, l2):
                w.bind("<Button-1>", lambda e, n=name: self.waehlen(n))
                w.bind("<Double-Button-1>", lambda e, n=name: self._doppel(n))
                w.bind("<MouseWheel>", self._rad)
            self.zeilen.append((name, f, l1, l2))
        self.inner.update_idletasks()
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))

    def _doppel(self, name):
        self.waehlen(name)
        if self.on_doppelklick:
            self.on_doppelklick(name)

    def waehlen(self, name):
        self.gewaehlt = name
        for i, (n, f, l1, l2) in enumerate(self.zeilen):
            bg = AUSWAHL if n == name else ZEBRA[i % 2]
            for w in (f, l1, l2):
                w.configure(bg=bg)

    def auswahl(self):
        return self.gewaehlt

    def sichtbar_machen(self, name):
        for n, f, _, _ in self.zeilen:
            if n == name:
                self.inner.update_idletasks()
                h = max(1, self.inner.winfo_height())
                self.canvas.yview_moveto(max(0.0, (f.winfo_y() - 10) / h))
                break


class FortschrittFenster(tk.Toplevel):
    """Grosses Fortschrittsfenster in der Bildschirmmitte."""

    def __init__(self, master, abbrechen, ziel_oeffnen):
        super().__init__(master)
        self.title("Fotos sortieren – Fortschritt")
        self.resizable(False, False)
        self.protocol("WM_DELETE_WINDOW", lambda: None)   # nicht wegklickbar, solange es läuft
        self.transient(master)
        self.abbrechen_cb, self.ziel_cb = abbrechen, ziel_oeffnen
        b, h = 820, 320
        sx, sy = self.winfo_screenwidth(), self.winfo_screenheight()
        self.geometry(f"{b}x{h}+{(sx - b) // 2}+{(sy - h) // 2}")
        self.configure(bg="white")

        kopf = tk.Frame(self, bg=KOPF_FARBE, padx=18, pady=10)
        kopf.pack(fill="x")
        self.l_titel = tk.Label(kopf, text="Fotos werden sortiert …", font=("Arial", GROESSE + 6, "bold"),
                                bg=KOPF_FARBE, fg="white", anchor="w")
        self.l_titel.pack(anchor="w")

        inhalt = tk.Frame(self, bg="white", padx=24, pady=16)
        inhalt.pack(fill="both", expand=True)
        st = ttk.Style()
        st.configure("Popup.Horizontal.TProgressbar", thickness=34)
        self.balken = ttk.Progressbar(inhalt, mode="indeterminate", style="Popup.Horizontal.TProgressbar", length=760)
        self.balken.pack(fill="x")
        self.balken.start(12)
        self.l_prozent = tk.Label(inhalt, text="Vorbereitung …", font=("Arial", GROESSE + 10, "bold"),
                                  bg="white", fg=KOPF_FARBE)
        self.l_prozent.pack(anchor="w", pady=(12, 0))
        self.l_zeile1 = tk.Label(inhalt, text="Dateien werden gesucht …", font=SCHRIFT, bg="white", anchor="w")
        self.l_zeile1.pack(anchor="w", pady=(4, 0))
        self.l_zeile2 = tk.Label(inhalt, text="", font=SCHRIFT_KLEIN, bg="white", fg=HINWEIS_FARBE, anchor="w",
                                 justify="left", wraplength=760)
        self.l_zeile2.pack(anchor="w")

        kn = tk.Frame(self, bg="white", padx=24, pady=12)
        kn.pack(fill="x")
        self.b_abbruch = ttk.Button(kn, text="■  Abbrechen", command=self._abbrechen)
        self.b_abbruch.pack(side="right")
        self.b_ziel = ttk.Button(kn, text="Zielordner öffnen", command=self.ziel_cb)
        self.b_liste = ttk.Button(kn, text="Zuordnungsliste anzeigen", command=lambda: master._liste_zeigen())
        self.b_schliessen = ttk.Button(kn, text="Schliessen", command=self.destroy)
        self.lift()
        self.focus_force()

    def _abbrechen(self):
        self.abbrechen_cb()
        self.l_zeile1.config(text="Wird abgebrochen … (aktuelle Datei wird noch fertig kopiert)")
        self.b_abbruch.config(state="disabled")

    def vorbereitung(self, text):
        self.l_zeile1.config(text=text)

    def fortschritt(self, n, t, verstrichen, rest):
        if str(self.balken.cget("mode")) != "determinate":
            self.balken.stop()
            self.balken.configure(mode="determinate")
        self.balken.configure(maximum=max(t, 1), value=n)
        self.l_prozent.config(text=f"{100.0 * n / max(t, 1):.0f} %")
        self.l_zeile1.config(text=f"Datei {n} von {t}")
        self.l_zeile2.config(text=f"verstrichen {App._zeit(verstrichen)}   ·   Restzeit ca. {App._zeit(rest)}")

    def fertig(self, text, gesamt, zusammenfassung=""):
        self.balken.stop()
        self.balken.configure(mode="determinate")
        if text == "Fertig.":
            self.balken.configure(value=self.balken.cget("maximum"))
            self.l_titel.config(text="Sortierung abgeschlossen")
            self.l_prozent.config(text="100 %")
        else:
            self.l_titel.config(text=text)
        self.l_zeile1.config(text=text + (f"   Gesamtdauer {App._zeit(gesamt)}" if gesamt else ""))
        self.l_zeile2.config(text=zusammenfassung)
        self.b_abbruch.pack_forget()
        self.b_schliessen.pack(side="right")
        self.b_ziel.pack(side="right", padx=(0, 8))
        self.b_liste.pack(side="right", padx=(0, 8))
        self.protocol("WM_DELETE_WINDOW", self.destroy)


# Spalten der Zuordnungsliste: (Protokollfeld, Überschrift, Breite Excel, Anteil Bildschirmbreite)
LISTEN_SPALTEN = [
    ("Nr", "Nr", 6, 3), ("Dateiname", "Dateiname", 28, 11), ("Quellordner", "Quellordner", 34, 10),
    ("Datum", "Aufnahmedatum", 19, 8), ("Datumsquelle", "Datum aus", 12, 5),
    ("Land", "Land", 14, 5), ("Ort", "Ort", 18, 6),
    ("Kategorie", "Thema", 22, 8), ("Sicherheit", "Sicherheit", 11, 4), ("Zweite_Wahl", "Zweite Wahl", 24, 8),
    ("Panorama", "Panorama", 10, 4),
    ("Doppelbild", "Doppel", 8, 4), ("Doppelbild_von", "Doppel von", 40, 10), ("Aehnlichkeit", "Ähnl. %", 8, 4),
    ("Ziel_Datum", "Zielordner Datum", 40, 9), ("Ziel_Ort", "Zielordner Ort", 40, 8), ("Ziel_Thema", "Zielordner Thema", 40, 8),
    ("Ziel_Doppel", "Zielordner Doppel", 40, 8),
]


def zeilen_aufbereiten(protokoll):
    """Protokollzeilen des Kerns in Listenzeilen (Dateiname/Quellordner getrennt) umwandeln."""
    out = []
    for z in protokoll:
        q = Path(str(z.get("Quelle", "")))
        d = dict(z)
        d["Dateiname"] = q.name
        d["Quellordner"] = q.parent.name
        out.append(d)
    return out


class ZuordnungsFenster(tk.Toplevel):
    """Tabelle mit der Zuordnung jedes Bildes, Excel-Export, Zurück-Knopf."""

    def __init__(self, master, protokoll, ziel: Path, testlauf: bool, programmordner: Path):
        super().__init__(master)
        self.title("Zuordnungsliste")
        self.zeilen = zeilen_aufbereiten(protokoll)
        self.ziel, self.testlauf, self.programmordner = ziel, testlauf, programmordner
        b, h = min(1500, self.winfo_screenwidth() - 80), min(900, self.winfo_screenheight() - 120)
        self.geometry(f"{b}x{h}+40+40")

        kopf = tk.Frame(self, bg=KOPF_FARBE, padx=16, pady=10)
        kopf.pack(fill="x")
        tk.Label(kopf, text="Zuordnungsliste", font=("Arial", GROESSE + 6, "bold"), bg=KOPF_FARBE,
                 fg="white").pack(side="left")
        tk.Label(kopf, text=f"{len(self.zeilen)} Dateien" + ("   ·   Testlauf" if testlauf else ""),
                 font=SCHRIFT, bg=KOPF_FARBE, fg="#DCE6F2").pack(side="left", padx=(20, 0), pady=(6, 0))

        kn = ttk.Frame(self, padding=(12, 8))
        kn.pack(fill="x")
        ttk.Button(kn, text="◀  Zurück", command=self.destroy).pack(side="left")
        ttk.Button(kn, text="Excelliste erstellen und öffnen", command=self.excel).pack(side="left", padx=(12, 0))
        self.l_info = ttk.Label(kn, text="Spaltenüberschrift anklicken = sortieren", style="Hinweis.TLabel")
        self.l_info.pack(side="left", padx=(20, 0))

        rahmen = ttk.Frame(self, padding=(12, 0, 12, 12))
        rahmen.pack(fill="both", expand=True)
        rahmen.rowconfigure(0, weight=1)
        rahmen.columnconfigure(0, weight=1)
        felder = [f for f, _, _, _ in LISTEN_SPALTEN]
        st = ttk.Style()
        st.configure("Liste.Treeview", font=SCHRIFT_TABELLE, rowheight=int((GROESSE - 3) * 2.1))
        st.configure("Liste.Treeview.Heading", font=("Arial", GROESSE - 3, "bold"))
        self.tv = ttk.Treeview(rahmen, columns=felder, show="headings", style="Liste.Treeview")
        for f, titel, breite, anteil in LISTEN_SPALTEN:
            self.tv.heading(f, text=titel, anchor="w", command=lambda c=f: self.sortieren(c, False))
            self.tv.column(f, width=anteil * 12, minwidth=30, anchor="w", stretch=True)
        sy = ttk.Scrollbar(rahmen, orient="vertical", command=self.tv.yview)
        self.tv.configure(yscrollcommand=sy.set)
        self.tv.grid(row=0, column=0, sticky="nsew")
        sy.grid(row=0, column=1, sticky="ns")
        self.tv.tag_configure("odd", background=ZEBRA[1])
        self.tv.bind("<Configure>", self._spalten_anpassen)
        self.fuellen(self.zeilen)
        try:
            if sys.platform.startswith("win"):
                self.state("zoomed")
        except Exception:
            pass

    def _spalten_anpassen(self, e):
        """Spaltenbreiten proportional auf die verfügbare Breite verteilen (kein horizontales Scrollen)."""
        gesamt = sum(a for _, _, _, a in LISTEN_SPALTEN)
        breite = max(300, e.width - 4)
        for f, _, _, anteil in LISTEN_SPALTEN:
            self.tv.column(f, width=int(breite * anteil / gesamt))

    @staticmethod
    def _kurz(f, v):
        # In der Tabelle nur den Zielordner zeigen (Dateiname steht schon in eigener Spalte); Excel behält den vollen Pfad
        if f == "Doppelbild_von" and v:
            return str(v).replace("/", "\\").rsplit("\\", 1)[-1]
        if f.startswith("Ziel_") and v and "\\" in str(v):
            teile = str(v).split(" | ")
            return " | ".join(t.rsplit("\\", 1)[0] for t in teile)
        return v

    def fuellen(self, zeilen):
        self.tv.delete(*self.tv.get_children())
        for i, z in enumerate(zeilen):
            self.tv.insert("", "end", values=[self._kurz(f, z.get(f, "")) for f, _, _, _ in LISTEN_SPALTEN],
                           tags=("odd",) if i % 2 else ())

    def sortieren(self, spalte, rueckwaerts):
        def key(z):
            v = z.get(spalte, "")
            try:
                return (0, float(v))
            except (TypeError, ValueError):
                return (1, str(v).lower())
        self.zeilen.sort(key=key, reverse=rueckwaerts)
        self.fuellen(self.zeilen)
        self.tv.heading(spalte, command=lambda: self.sortieren(spalte, not rueckwaerts))

    def excel(self):
        try:
            import openpyxl
        except ImportError:
            if not messagebox.askyesno("Modul fehlt", "Für Excel-Dateien wird das Modul «openpyxl» benötigt.\n"
                                                      "Jetzt installieren (Internet nötig)?", parent=self):
                return
            r = subprocess.run([sys.executable.replace("pythonw.exe", "python.exe"), "-m", "pip", "install",
                                "openpyxl"], capture_output=True, text=True)
            try:
                import openpyxl  # noqa
            except ImportError:
                messagebox.showerror("Installation fehlgeschlagen", r.stderr[-1500:] or r.stdout[-1500:], parent=self)
                return
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Zuordnung"
        ws.append([titel for _, titel, _, _ in LISTEN_SPALTEN])
        for z in self.zeilen:
            zeile = []
            for f, _, _, _ in LISTEN_SPALTEN:
                v = z.get(f, "")
                if f in ("Nr",):
                    try:
                        v = int(v)
                    except (TypeError, ValueError):
                        pass
                elif f == "Sicherheit":
                    try:
                        v = float(v)
                    except (TypeError, ValueError):
                        pass
                zeile.append(v)
            ws.append(zeile)
        kopf_fill = PatternFill("solid", fgColor="1F4E79")
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = kopf_fill
            c.alignment = Alignment(vertical="center")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial")
        for i, (_, _, breite, _) in enumerate(LISTEN_SPALTEN, 1):
            ws.column_dimensions[get_column_letter(i)].width = breite
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        # Zusammenfassung
        zs = wb.create_sheet("Zusammenfassung")
        zs.append(["Kennzahl", "Anzahl"])
        zs["A1"].font = zs["B1"].font = Font(name="Arial", bold=True)
        zaehl = {}
        for z in self.zeilen:
            for f, titel in (("Kategorie", "Thema"), ("Ort", "Ort"), ("Datumsquelle", "Datum aus"), ("Doppelbild", "Doppelbild")):
                v = z.get(f, "")
                if v:
                    zaehl[f"{titel}: {v}"] = zaehl.get(f"{titel}: {v}", 0) + 1
        zs.append(["Dateien total", len(self.zeilen)])
        for k in sorted(zaehl):
            zs.append([k, zaehl[k]])
        zs.column_dimensions["A"].width = 45
        zs.column_dimensions["B"].width = 10

        basis = self.programmordner if self.testlauf else self.ziel
        try:
            basis.mkdir(parents=True, exist_ok=True)
        except Exception:
            basis = self.programmordner
        pfad = basis / f"Zuordnungsliste_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        try:
            wb.save(pfad)
        except Exception as ex:
            messagebox.showerror("Speichern fehlgeschlagen", str(ex), parent=self)
            return
        self.l_info.config(text=f"Gespeichert: {pfad}")
        try:
            if sys.platform.startswith("win"):
                os.startfile(pfad)  # type: ignore[attr-defined]
        except Exception:
            pass


class ThemenFenster(tk.Toplevel):
    """Vollbildfenster zum Bearbeiten der Themen (Tabelle gross, Erläuterungen mit Abstand)."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        app.themen_fenster = self
        self.title("Themen bearbeiten")
        try:
            if sys.platform.startswith("win"):
                self.state("zoomed")
            else:
                self.geometry(f"{self.winfo_screenwidth()}x{self.winfo_screenheight() - 60}+0+0")
        except Exception:
            pass
        self.transient(app)

        kopf = tk.Frame(self, bg=KOPF_FARBE, padx=16, pady=10)
        kopf.pack(fill="x")
        tk.Label(kopf, text="Themen bearbeiten", font=("Arial", GROESSE + 6, "bold"), bg=KOPF_FARBE,
                 fg="white").pack(side="left")
        tk.Label(kopf, text="Ein Thema = ein Ordner in nach_Thema.  Die Beschreibungen sagen dem KI-Modell, "
                            "welche Bilder hineingehören.", font=SCHRIFT, bg=KOPF_FARBE, fg="#DCE6F2").pack(
            side="left", padx=(24, 0), pady=(6, 0))

        kn = ttk.Frame(self, padding=(16, 10, 16, 4))
        kn.pack(fill="x")
        ttk.Button(kn, text="◀  Zurück (übernehmen)", command=self.schliessen).pack(side="left")
        ttk.Button(kn, text="Ergänzen / Neu ...", command=app._kat_neu).pack(side="left", padx=(16, 0))
        ttk.Button(kn, text="Bearbeiten ...", command=app._kat_bearbeiten).pack(side="left", padx=(8, 0))
        ttk.Button(kn, text="Löschen", command=app._kat_loeschen).pack(side="left", padx=(8, 0))
        ttk.Button(kn, text="Standard laden", command=app._kat_standard).pack(side="left", padx=(8, 0))
        ttk.Label(kn, text="Doppelklick auf eine Zeile = bearbeiten", style="Erlaeuterung.TLabel").pack(side="left", padx=(20, 0))

        inhalt = ttk.Frame(self, padding=(16, 4, 16, 12))
        inhalt.pack(fill="both", expand=True)
        inhalt.rowconfigure(0, weight=1)
        inhalt.columnconfigure(0, weight=1)
        self.liste = ThemenTabelle(inhalt, on_doppelklick=lambda n: app._dialog_ausfuehren(n), spalte1=300)
        self.liste.grid(row=0, column=0, sticky="nsew")
        self.liste.setzen(app.kategorien)

        # Anzeige-Hinweis (Übersetzung) direkt unter der Tabelle
        self.l_anzeige = ttk.Label(inhalt, text="", style="Erlaeuterung.TLabel", wraplength=1500, justify="left")
        self.l_anzeige.grid(row=1, column=0, sticky="w", pady=(8, 0))

        # Modell + Schwelle (mit grösserem Abstand)
        unten = ttk.Frame(inhalt)
        unten.grid(row=2, column=0, sticky="ew", pady=(24, 0))
        ttk.Label(unten, text="KI-Modell:", font=SCHRIFT_FETT).pack(side="left")
        cb = ttk.Combobox(unten, textvariable=app.v_modell, state="readonly", width=16, values=list(kern.MODELLE))
        cb.pack(side="left", padx=(6, 12))
        cb.bind("<<ComboboxSelected>>", lambda e: (app._modell_info(), self.anzeige_aktualisieren()))
        self.l_modell = ttk.Label(unten, text=kern.MODELLE[app.v_modell.get()][2], style="Erlaeuterung.TLabel")
        self.l_modell.pack(side="left")
        ttk.Label(unten, text="      Mindest-Sicherheit:", font=SCHRIFT_FETT).pack(side="left")
        ttk.Spinbox(unten, from_=0.1, to=0.95, increment=0.05, textvariable=app.v_schwelle, width=6,
                    format="%.2f", command=app._themen_kurz_aktualisieren).pack(side="left", padx=6)
        ttk.Label(unten, text="(darunter → Ordner _unsicher)", style="Erlaeuterung.TLabel").pack(side="left")

        # Erläuterungen mit Abstand
        erl = ttk.Frame(inhalt)
        erl.grid(row=3, column=0, sticky="ew", pady=(20, 0))
        for zeile in (
            "Richtwerte Mindest-Sicherheit:   ≥ 0.70 = ziemlich sicher (selten falsch)   ·   0.45–0.70 = wahrscheinlich "
            "richtig, Grenzfälle möglich (zweite Wahl im Protokoll beachten)   ·   < 0.45 = unsicher → Ordner _unsicher",
            "Die Werte sind Anteile über alle Themen (Summe 100 %): je mehr Themen erfasst sind, desto tiefer fallen die "
            "Werte im Schnitt aus. Bei über 12 Themen die Schwelle eher auf 0.35–0.40 senken.",
            "Mehrere Formulierungen pro Thema machen die Erkennung robuster. Das Modell erkennt Motive, keine Orte oder "
            "Personen – Ortsangaben in Beschreibungen bringen nichts.",
            "Englische Modelle (englisch, englisch_gross): deutsche Beschreibungen werden automatisch ins Englische "
            "übersetzt (lokal, einmalig ca. 300 MB Download); bereits englische Texte bleiben unverändert.",
        ):
            ttk.Label(erl, text=zeile, style="Erlaeuterung.TLabel", wraplength=1500, justify="left").pack(
                anchor="w", pady=(0, 6))

        self.protocol("WM_DELETE_WINDOW", self.schliessen)
        self._meldungen = queue.Queue()
        self._uebersetzt = None
        self.anzeige_aktualisieren()

    # ---- Anzeige der Themen: deutsch oder (bei englischem Modell) automatisch übersetzt
    def anzeige_aktualisieren(self):
        if self.app.v_modell.get() == "mehrsprachig":
            self._uebersetzt = None
            self.liste.setzen(self.app.kategorien)
            self.l_anzeige.config(text="Anzeige: Beschreibungen auf Deutsch (mehrsprachiges Modell).")
            return
        self.l_anzeige.config(text="Englisches Modell gewählt – Beschreibungen werden übersetzt … "
                                   "(beim ersten Mal Download des Übersetzungsmodells, ca. 300 MB, bitte warten)")
        kategorien = {k: list(v) for k, v in self.app.kategorien.items()}

        def arbeit():
            try:
                ueb = kern.Uebersetzer(PROGRAMMORDNER / "uebersetzungen.json",
                                       lambda t: self._meldungen.put(("log", t)))
                ergebnis = {k: ueb.uebersetzen(v) for k, v in kategorien.items()}
                self._meldungen.put(("fertig", (ergebnis, ueb.fehler)))
            except Exception as ex:
                self._meldungen.put(("fertig", (None, f"{type(ex).__name__}: {ex}")))

        threading.Thread(target=arbeit, daemon=True).start()
        self.after(200, self._abholen)

    def _abholen(self):
        if not self.winfo_exists():
            return
        try:
            while True:
                art, wert = self._meldungen.get_nowait()
                if art == "log":
                    self.l_anzeige.config(text=wert)
                elif art == "fertig":
                    ergebnis, fehler = wert
                    if ergebnis and not fehler:
                        self._uebersetzt = ergebnis
                        self.liste.setzen(ergebnis)
                        self.l_anzeige.config(
                            text="Anzeige: englische Übersetzung, wie sie das KI-Modell verwendet. Gespeichert und "
                                 "bearbeitet werden die Themen weiterhin auf Deutsch (Bearbeiten-Dialog zeigt den "
                                 "deutschen Text); Ordnernamen bleiben unverändert.")
                    else:
                        self.liste.setzen(self.app.kategorien)
                        self.l_anzeige.config(text=f"Übersetzung nicht möglich ({fehler}). Anzeige auf Deutsch; "
                                                   "beim Lauf wird der deutsche Text verwendet.")
                    return
        except queue.Empty:
            pass
        self.after(200, self._abholen)

    def schliessen(self):
        self.app._themen_kurz_aktualisieren()
        self.app._einstellungen_uebernehmen()
        self.app.themen_fenster = None
        self.destroy()


class App(tk.Tk):
    def __init__(self):
        # Scharfe Darstellung auf hochauflösenden Bildschirmen (Windows)
        if sys.platform.startswith("win"):
            try:
                import ctypes
                ctypes.windll.shcore.SetProcessDpiAwareness(1)
            except Exception:
                pass
        super().__init__()
        # Tk-Skalierung an die tatsächliche Bildschirm-DPI anpassen (sonst zu kleine Schrift bei 150–250 %)
        if sys.platform.startswith("win"):
            try:
                import ctypes
                dpi = ctypes.windll.user32.GetDpiForSystem()
                self.tk.call("tk", "scaling", dpi / 72.0)
            except Exception:
                pass
        self.title(f"Fotosortierprogramm  –  Version {VERSION}")
        self.minsize(1100, 800)
        self.geometry("1500x1000")
        # Fenster beim Start bildschirmfüllend (Windows: maximiert, sonst volle Bildschirmgrösse)
        try:
            if sys.platform.startswith("win"):
                self.state("zoomed")
            else:
                self.geometry(f"{self.winfo_screenwidth()}x{self.winfo_screenheight() - 60}+0+0")
        except Exception:
            pass
        # Schrift: Arial, gut lesbar, für alle Elemente
        self.option_add("*Font", SCHRIFT)
        self.option_add("*Listbox.Font", SCHRIFT)
        self.option_add("*Text.Font", SCHRIFT)
        self.option_add("*Entry.Font", SCHRIFT)
        self.option_add("*TCombobox*Listbox.Font", SCHRIFT)
        try:
            if ICON.exists():
                self.iconbitmap(default=str(ICON))
        except Exception:
            pass
        st = ttk.Style()
        try:
            st.theme_use("vista" if sys.platform.startswith("win") else "clam")
        except Exception:
            pass
        st.configure(".", font=SCHRIFT)
        st.configure("TLabelframe.Label", font=SCHRIFT_FETT)
        st.configure("TButton", font=SCHRIFT, padding=(8, 4))
        st.configure("TCheckbutton", font=SCHRIFT, padding=(0, 3))
        st.configure("TEntry", font=SCHRIFT)
        st.configure("Hinweis.TLabel", font=SCHRIFT_KLEIN, foreground=HINWEIS_FARBE)
        st.configure("Erlaeuterung.TLabel", font=SCHRIFT_ERLAEUTERUNG, foreground=HINWEIS_FARBE)
        st.configure("Treeview", font=SCHRIFT, rowheight=int(GROESSE * 2.2))
        st.configure("Treeview.Heading", font=SCHRIFT_FETT)

        self.e = kern.einstellungen_laden(EINSTELLUNGSDATEI)
        self.kategorien = dict(self.e.get("kategorien") or kern.STANDARD_KATEGORIEN)
        self.meldungen = queue.Queue()
        self.abbruch = threading.Event()
        self.thread = None

        self._aufbau()
        self.after(150, self._meldungen_abholen)
        self.protocol("WM_DELETE_WINDOW", self._schliessen)
        # Update-Prüfung im Hintergrund (höchstens einmal pro Tag)
        self.after(1500, self._update_pruefen)
        # Erster Start: Quellordner fehlt -> Abschnitt 1 aufklappen und Hinweis geben
        if not str(self.e.get("quelle", "")).strip():
            self.f_ordner.setzen(True)
            self.l_status.config(text="Erster Start: Bitte unter «1. Ordner» den Quellordner mit den Fotos wählen. "
                                      "Ziel- und Doppelbilder-Ordner sind vorbelegt und können geändert werden.")

    # ------------------------------------------------------------------ Aufbau
    def _haupt_breite(self, e):
        self.haupt_canvas.itemconfigure(self._haupt_fenster, width=e.width)

    def _haupt_rad(self, e):
        # Nur das Hauptfenster scrollen (nicht Themen-/Listenfenster)
        try:
            if e.widget.winfo_toplevel() is not self:
                return
        except Exception:
            pass
        if self.haupt_canvas.bbox("all") and self.haupt_canvas.bbox("all")[3] > self.haupt_canvas.winfo_height():
            self.haupt_canvas.yview_scroll(int(-e.delta / 120), "units")

    def _abschnitt(self, parent, titel, row, weight=0, hinweis=None, offen=True, klappbar=True, hell=None):
        """Blauer Titelbalken (+ Hinweiszeile) + Inhaltsrahmen; Klick auf den Balken klappt auf/zu."""
        rahmen = ttk.Frame(parent)
        rahmen.grid(row=row, column=0, sticky="nsew", pady=(34 if row else 10, 0))
        rahmen.columnconfigure(0, weight=1)
        rahmen.rowconfigure(2, weight=1)
        farbe = BALKEN_HELL if (klappbar if hell is None else hell) else BALKEN_FARBE
        balken = tk.Frame(rahmen, bg=farbe, cursor="hand2" if klappbar else "")
        balken.grid(row=0, column=0, sticky="ew")
        balken.columnconfigure(1, weight=1)
        pfeil = tk.Label(balken, text="▾" if offen else "▸", font=SCHRIFT_ABSCHNITT, bg=farbe,
                         fg=BALKEN_TEXT, width=2, anchor="w", padx=10, pady=7)
        pfeil.grid(row=0, column=0, sticky="w")
        if not klappbar:
            pfeil.configure(text="")
        l_titel = tk.Label(balken, text=titel, font=SCHRIFT_ABSCHNITT, bg=farbe, fg=BALKEN_TEXT,
                           anchor="w", pady=7)
        l_titel.grid(row=0, column=1, sticky="w")
        l_kurz = tk.Label(balken, text="", font=SCHRIFT_KLEIN, bg=farbe, fg="#EAF1FA" if klappbar else "#C9D8EA",
                          anchor="e", padx=14)
        l_kurz.grid(row=0, column=2, sticky="e")
        info = None
        if hinweis:
            info = ttk.Label(rahmen, text=hinweis, style="Hinweis.TLabel", wraplength=1400, justify="left")
            info.grid(row=1, column=0, sticky="w", padx=14, pady=(4, 0))
        inhalt = ttk.Frame(rahmen, padding=(14, 8, 14, 6))
        inhalt.grid(row=2, column=0, sticky="nsew")
        zustand = {"offen": offen}

        def setzen(neu):
            zustand["offen"] = neu
            pfeil.configure(text="▾" if neu else "▸")
            if neu:
                if info is not None:
                    info.grid()
                inhalt.grid()
                if weight:
                    parent.rowconfigure(row, weight=weight)
            else:
                if info is not None:
                    info.grid_remove()
                inhalt.grid_remove()
                parent.rowconfigure(row, weight=0)

        def umschalten(_=None):
            setzen(not zustand["offen"])

        if klappbar:
            for w in (balken, pfeil, l_titel, l_kurz):
                w.bind("<Button-1>", umschalten)
        inhalt.kurz = l_kurz          # Kurzinfo im Balken (z. B. gewählte Ordner)
        inhalt.setzen = setzen
        if not offen:
            setzen(False)
        if weight and offen:
            parent.rowconfigure(row, weight=weight)
        return inhalt

    def _kurzinfos(self, *_):
        """Kurzinfo in den zugeklappten Balken aktualisieren."""
        try:
            q, z = self.v_quelle.get().strip(), self.v_ziel.get().strip()
            self.f_ordner.kurz.config(text=(f"Quelle: {Path(q).name or q}   →   Ziel: {Path(z).name or z}"
                                            if q or z else "noch keine Ordner gewählt"))
            arten = [n for n, v in (("Datum", self.v_datum), ("Ort", self.v_ort), ("Thema", self.v_thema),
                                    ("Panorama", self.v_pano), ("Doppelbilder", self.v_doppel)) if v.get()]
            self.f_art.kurz.config(text=", ".join(arten) if arten else "nichts gewählt")
            self._themen_kurz_aktualisieren()
        except Exception:
            pass

    def _aufbau(self):
        # Kopfbereich (dunkelblau, volle Breite)
        kopf = tk.Frame(self, bg=KOPF_FARBE, padx=16, pady=12)
        kopf.pack(fill="x")
        tk.Label(kopf, text="Fotosortierprogramm", font=SCHRIFT_TITEL, bg=KOPF_FARBE, fg="white",
                 anchor="w").pack(anchor="w")
        tk.Label(kopf, text="Fotos nach Aufnahmedatum, Aufnahmeort und Bildthema sortieren – die Originale bleiben unverändert",
                 font=SCHRIFT_UNTERTITEL, bg=KOPF_FARBE, fg="#DCE6F2", anchor="w").pack(anchor="w", pady=(2, 0))
        vz = tk.Frame(kopf, bg=KOPF_FARBE)
        vz.pack(anchor="w", pady=(4, 0), fill="x")
        tk.Label(vz, text=f"Version {VERSION}", font=("Arial", GROESSE - 1, "bold"), bg=KOPF_FARBE,
                 fg="#9CC3EA", anchor="w").pack(side="left")
        self.b_update = tk.Button(vz, text="Nach Updates suchen", font=SCHRIFT_ERLAEUTERUNG, bg="#2E5C8E", fg="white",
                                  activebackground="#3D6EA5", activeforeground="white", relief="flat", padx=10, pady=2,
                                  cursor="hand2", command=lambda: self._update_pruefen(manuell=True))
        self.b_update.pack(side="left", padx=(24, 0))
        self.l_update = tk.Label(vz, text="", font=SCHRIFT_ERLAEUTERUNG, bg=KOPF_FARBE, fg="#DCE6F2", anchor="w")
        self.l_update.pack(side="left", padx=(12, 0))

        # Scrollbarer Hauptbereich (bei kleinen Bildschirmen)
        aussen = ttk.Frame(self)
        aussen.pack(fill="both", expand=True)
        self.haupt_canvas = tk.Canvas(aussen, highlightthickness=0)
        haupt_sb = ttk.Scrollbar(aussen, orient="vertical", command=self.haupt_canvas.yview)
        self.haupt_canvas.configure(yscrollcommand=haupt_sb.set)
        haupt_sb.pack(side="right", fill="y")
        self.haupt_canvas.pack(side="left", fill="both", expand=True)
        haupt = ttk.Frame(self.haupt_canvas, padding=(14, 6, 14, 10))
        self._haupt_fenster = self.haupt_canvas.create_window((0, 0), window=haupt, anchor="nw")
        haupt.bind("<Configure>", lambda e: self.haupt_canvas.configure(scrollregion=self.haupt_canvas.bbox("all")))
        self.haupt_canvas.bind("<Configure>", self._haupt_breite)
        self.bind_all("<MouseWheel>", self._haupt_rad)
        haupt.columnconfigure(0, weight=1)

        # --- Ordner
        self.f_ordner = f_ordner = self._abschnitt(haupt, "1.  Ordner", 0, offen=False, hinweis="Quellordner mit den Fotos (alle Unterordner werden durchsucht) und Zielordner für die sortierten Kopien.")
        f_ordner.columnconfigure(1, weight=1)
        self.v_quelle = tk.StringVar(value=self.e["quelle"])
        self.v_ziel = tk.StringVar(value=self.e["ziel"])
        ttk.Label(f_ordner, text="Fotos (Quelle):").grid(row=0, column=0, sticky="w")
        ttk.Entry(f_ordner, textvariable=self.v_quelle).grid(row=0, column=1, sticky="ew", padx=6)
        ttk.Button(f_ordner, text="Wählen ...", command=lambda: self._ordner_waehlen(self.v_quelle)).grid(row=0, column=2)
        ttk.Label(f_ordner, text="Ziel (sortiert):").grid(row=1, column=0, sticky="w")
        ttk.Entry(f_ordner, textvariable=self.v_ziel).grid(row=1, column=1, sticky="ew", padx=6)
        ttk.Button(f_ordner, text="Wählen ...", command=lambda: self._ordner_waehlen(self.v_ziel)).grid(row=1, column=2)
        ttk.Label(f_ordner, text="Unterordner ausschliessen:").grid(row=2, column=0, sticky="w", pady=(6, 0))
        self.v_ausschluss = tk.StringVar(value=self.e.get("ausschluss", ""))
        ttk.Entry(f_ordner, textvariable=self.v_ausschluss).grid(row=2, column=1, sticky="ew", padx=6, pady=(6, 0))
        ttk.Label(f_ordner, text="Doppelbilder-Ordner:").grid(row=3, column=0, sticky="w", pady=(6, 0))
        self.v_doppel_ordner = tk.StringVar(value=self.e.get("doppel_ordner", ""))
        ttk.Entry(f_ordner, textvariable=self.v_doppel_ordner).grid(row=3, column=1, sticky="ew", padx=6, pady=(6, 0))
        ttk.Button(f_ordner, text="Wählen ...", command=lambda: self._ordner_waehlen(self.v_doppel_ordner)).grid(
            row=3, column=2, pady=(6, 0))
        ttk.Label(f_ordner, style="Hinweis.TLabel", wraplength=1300, justify="left",
                  text="Doppelbilder-Ordner: Dorthin werden doppelte oder sehr ähnliche Bilder kopiert (leer = "
                       "Zielordner\\Aussortierte_Doppelbilder). Ausschluss: Ordnernamen durch Komma getrennt "
                       "(z. B. technische Ordner oder Originale aus einem Bildbearbeitungsprogramm). "
                       "Ziel nicht innerhalb der Quelle wählen; bei synchronisierten Ordnern das Ziel ausserhalb "
                       "des Sync-Bereichs legen.").grid(row=4, column=0, columnspan=3, sticky="w", pady=(4, 0))

        # --- Sortierarten
        self.f_art = f_art = self._abschnitt(haupt, "2.  Wie sortieren?", 1, offen=False, hinweis="Mehrere Sortierarten sind gleichzeitig möglich; jede erzeugt einen eigenen Unterordner im Ziel.")
        self.v_datum = tk.BooleanVar(value=self.e["nach_datum"])
        self.v_ort = tk.BooleanVar(value=self.e["nach_ort"])
        self.v_thema = tk.BooleanVar(value=self.e["nach_thema"])
        self.v_pano = tk.BooleanVar(value=self.e["panorama"])
        self.v_struktur = tk.StringVar(value=self.e["datum_struktur"])
        ttk.Checkbutton(f_art, text="Nach Aufnahmedatum  →  nach_Datum\\", variable=self.v_datum).grid(row=0, column=0, sticky="w")
        ttk.Label(f_art, text="Struktur:").grid(row=0, column=1, sticky="e", padx=(20, 4))
        ttk.Combobox(f_art, textvariable=self.v_struktur, state="readonly", width=22,
                     values=["Jahr/Jahr-Monat", "Jahr/Jahr-Monat-Tag", "Jahr-Monat"]).grid(row=0, column=2, sticky="w")
        ttk.Checkbutton(f_art, text="Nach Ort (GPS, Internet nötig)  →  nach_Ort\\Land\\Ort\\", variable=self.v_ort).grid(
            row=1, column=0, sticky="w")
        ttk.Checkbutton(f_art, text="Nach Thema (Bildinhalt, lokales KI-Modell)  →  nach_Thema\\Kategorie\\",
                        variable=self.v_thema, command=self._thema_umschalten).grid(row=2, column=0, sticky="w")
        ttk.Checkbutton(f_art, text="Panoramen zusätzlich sammeln  →  nach_Thema\\Panorama\\", variable=self.v_pano).grid(
            row=3, column=0, sticky="w")
        self.v_doppel = tk.BooleanVar(value=bool(self.e.get("doppel_erkennen", True)))
        self.v_doppel_proz = tk.IntVar(value=int(self.e.get("doppel_aehnlichkeit", 90)))
        ttk.Checkbutton(f_art, text="Doppelte und sehr ähnliche Bilder aussortieren  →  Doppelbilder-Ordner",
                        variable=self.v_doppel).grid(row=4, column=0, sticky="w")
        ttk.Label(f_art, text="ab Ähnlichkeit:").grid(row=4, column=1, sticky="e", padx=(20, 4))
        fr_d = ttk.Frame(f_art)
        fr_d.grid(row=4, column=2, sticky="w")
        ttk.Spinbox(fr_d, from_=70, to=100, increment=1, textvariable=self.v_doppel_proz, width=5).pack(side="left")
        ttk.Label(fr_d, text="%   (100 = nur identische Bilder; 90 = auch leicht veränderte Kopien)",
                  style="Hinweis.TLabel").pack(side="left", padx=(6, 0))

        # --- Themen (Bearbeitung in eigenem Vollbildfenster)
        self.f_thema = self._abschnitt(haupt, "3.  Themen (frei ergänzbar)", 2, klappbar=False, hell=True,
                                       hinweis="Nur für die Sortierung nach Thema. Ein Thema = ein Ordner; die "
                                               "Beschreibungen sagen dem KI-Modell, was hineingehört.")
        self.v_modell = tk.StringVar(value=self.e["modell"])
        self.v_schwelle = tk.DoubleVar(value=float(self.e["schwelle"]))
        zeile = ttk.Frame(self.f_thema)
        zeile.grid(row=0, column=0, sticky="ew")
        self.b_themen = ttk.Button(zeile, text="Themen bearbeiten (Vollbild) ...", command=self._themen_fenster)
        self.b_themen.pack(side="left")
        self.l_themen_kurz = ttk.Label(zeile, text="", style="Hinweis.TLabel", wraplength=1100, justify="left")
        self.l_themen_kurz.pack(side="left", padx=(16, 0))
        self._themen_kurz_aktualisieren()
        self._thema_umschalten()

        # --- Ausführung
        f_run = self._abschnitt(haupt, "4.  Ausführen", 3, klappbar=False, hinweis="Die Fotos werden kopiert – die Originale verbleiben unverändert in den ursprünglichen Ordnern. Zuerst Testlauf (nur Protokoll), Protokoll prüfen, dann echten Lauf starten.")
        f_run.columnconfigure(3, weight=1)
        self.v_test = tk.BooleanVar(value=self.e["testlaufe"])
        self.v_move = tk.BooleanVar(value=self.e["verschieben"])
        ttk.Checkbutton(f_run, text="Testlauf (nur Protokoll, nichts kopieren)", variable=self.v_test).grid(
            row=0, column=0, sticky="w")
        ttk.Checkbutton(f_run, text="Originale nach nach_Datum verschieben statt kopieren (Vorsicht)",
                        variable=self.v_move).grid(row=0, column=1, sticky="w", padx=(16, 0))
        self.b_start = ttk.Button(f_run, text="▶  Start", command=self._start, width=14)
        self.b_start.grid(row=1, column=0, sticky="w", pady=(8, 0))
        self.b_stop = ttk.Button(f_run, text="■  Abbrechen", command=self._abbrechen, width=14, state="disabled")
        self.b_stop.grid(row=1, column=1, sticky="w", pady=(8, 0), padx=(8, 0))
        ttk.Button(f_run, text="Zielordner öffnen", command=self._ziel_oeffnen).grid(row=1, column=2, sticky="w",
                                                                                     pady=(8, 0), padx=(8, 0))
        # Fortschrittsanzeige (eigene Zeile, volle Breite)
        st = ttk.Style()
        st.configure("Gross.Horizontal.TProgressbar", thickness=26)
        self.fortschritt = ttk.Progressbar(f_run, mode="determinate", style="Gross.Horizontal.TProgressbar")
        self.fortschritt.grid(row=2, column=0, columnspan=4, sticky="ew", pady=(12, 0))
        self.l_status = ttk.Label(f_run, text="Bereit.", font=SCHRIFT_FETT)
        self.l_status.grid(row=3, column=0, columnspan=4, sticky="w", pady=(6, 0))
        self.l_status2 = ttk.Label(f_run, text="", style="Hinweis.TLabel")
        self.l_status2.grid(row=4, column=0, columnspan=4, sticky="w")
        self._startzeit = None

        # --- Protokoll
        f_log = self._abschnitt(haupt, "Protokoll", 4, weight=1, klappbar=False)
        self.f_log = f_log
        f_log.columnconfigure(0, weight=1)
        f_log.rowconfigure(0, weight=1)
        kn_log = ttk.Frame(f_log)
        kn_log.grid(row=1, column=0, sticky="w", pady=(6, 0))
        self.b_liste = ttk.Button(kn_log, text="Zuordnungsliste anzeigen", command=self._liste_zeigen, state="disabled")
        self.b_liste.pack(side="left")
        self.b_excel = ttk.Button(kn_log, text="Excelliste erstellen", command=self._excel_direkt, state="disabled")
        self.b_excel.pack(side="left", padx=(8, 0))
        ttk.Label(kn_log, text="(nach einem Lauf verfügbar)", style="Hinweis.TLabel").pack(side="left", padx=(10, 0))
        self.v_logtext = tk.BooleanVar(value=False)
        ttk.Checkbutton(kn_log, text="Protokolltext anzeigen", variable=self.v_logtext,
                        command=self._logtext_umschalten).pack(side="left", padx=(24, 0))
        self.log = scrolledtext.ScrolledText(f_log, height=8, wrap="word", state="disabled",
                                             font=SCHRIFT_MONO if sys.platform.startswith("win") else SCHRIFT)
        self.log.grid(row=0, column=0, sticky="nsew")
        self.log.grid_remove()
        f_log.rowconfigure(0, weight=0)
        # Kurzinfos in den Balken pflegen
        for v in (self.v_quelle, self.v_ziel, self.v_datum, self.v_ort, self.v_thema, self.v_pano,
                  self.v_doppel, self.v_modell, self.v_schwelle):
            v.trace_add("write", self._kurzinfos)
        self._kurzinfos()

    # ------------------------------------------------------------------ Kategorien
    def _kat_liste_aufbauen(self):
        if getattr(self, "themen_fenster", None) and self.themen_fenster.winfo_exists():
            self.themen_fenster.anzeige_aktualisieren()
        self._themen_kurz_aktualisieren()

    def _kat_gewaehlt(self):
        if getattr(self, "themen_fenster", None) and self.themen_fenster.winfo_exists():
            return self.themen_fenster.liste.auswahl()
        return None

    def _dialog_ausfuehren(self, vorwahl=None):
        eltern = self.themen_fenster if getattr(self, "themen_fenster", None) and self.themen_fenster.winfo_exists() else self
        d = KategorienDialog(eltern, self.kategorien, vorwahl, self.v_modell.get() == "mehrsprachig")
        self.wait_window(d)
        if d.ergebnis:
            alt, neu, b = d.ergebnis
            if alt and alt != neu and alt in self.kategorien:
                del self.kategorien[alt]
            self.kategorien[neu] = b
            self._kat_liste_aufbauen()
            self._kurzinfos()
            if eltern is not self:
                eltern.liste.waehlen(neu)
                eltern.liste.sichtbar_machen(neu)
            self.v_thema.set(True)
            self._thema_umschalten()

    def _kat_neu(self):
        # Auswahlfenster: bestehendes Thema ergänzen oder neues anlegen
        self._dialog_ausfuehren(None)

    def _kat_bearbeiten(self):
        name = self._kat_gewaehlt()
        if not name:
            messagebox.showinfo("Hinweis", "Bitte zuerst ein Thema in der Liste anklicken – oder «Hinzufügen» "
                                           "und dort das Thema auswählen.")
            return
        self._dialog_ausfuehren(name)

    def _kat_loeschen(self):
        name = self._kat_gewaehlt()
        if name and messagebox.askyesno("Löschen", f"Kategorie «{name}» löschen?"):
            del self.kategorien[name]
            self._kat_liste_aufbauen()

    def _kat_standard(self):
        if messagebox.askyesno("Standard", "Eigene Kategorien verwerfen und Standardliste laden?"):
            self.kategorien = json.loads(json.dumps(kern.STANDARD_KATEGORIEN))
            self._kat_liste_aufbauen()

    def _modell_info(self):
        if getattr(self, "themen_fenster", None) and self.themen_fenster.winfo_exists():
            self.themen_fenster.l_modell.config(text=kern.MODELLE[self.v_modell.get()][2])
        self._themen_kurz_aktualisieren()

    def _thema_umschalten(self):
        self.b_themen.configure(state="normal" if self.v_thema.get() else "disabled")
        self._themen_kurz_aktualisieren()

    def _themen_kurz_aktualisieren(self):
        namen = list(self.kategorien.keys())
        text = (f"{len(namen)} Themen   ·   Modell {self.v_modell.get()}   ·   Mindest-Sicherheit "
                f"{float(self.v_schwelle.get()):.2f}\n" + ", ".join(namen[:10]) + (" …" if len(namen) > 10 else ""))
        if not self.v_thema.get():
            text = "Sortierung nach Thema ist in Abschnitt 2 nicht angehakt.   " + text
        self.l_themen_kurz.configure(text=text)

    def _themen_fenster(self):
        ThemenFenster(self)

    # ------------------------------------------------------------------ Aktionen
    def _ordner_waehlen(self, var):
        p = filedialog.askdirectory(initialdir=var.get() or str(Path.home()))
        if p:
            var.set(os.path.normpath(p))

    def _einstellungen_uebernehmen(self):
        self.e.update(quelle=self.v_quelle.get().strip(), ziel=self.v_ziel.get().strip(),
                      nach_datum=self.v_datum.get(), nach_ort=self.v_ort.get(), nach_thema=self.v_thema.get(),
                      panorama=self.v_pano.get(), datum_struktur=self.v_struktur.get(),
                      verschieben=self.v_move.get(), testlaufe=self.v_test.get(),
                      modell=self.v_modell.get(), schwelle=float(self.v_schwelle.get()),
                      kategorien=self.kategorien, ausschluss=self.v_ausschluss.get().strip(),
                      doppel_erkennen=self.v_doppel.get(), doppel_aehnlichkeit=int(self.v_doppel_proz.get()),
                      doppel_ordner=self.v_doppel_ordner.get().strip())
        try:
            kern.einstellungen_speichern(EINSTELLUNGSDATEI, self.e)
        except Exception:
            pass

    def _schreiben(self, text):
        self.meldungen.put(("log", text))

    def _meldungen_abholen(self):
        try:
            while True:
                art, wert = self.meldungen.get_nowait()
                if art == "log":
                    self.log.configure(state="normal")
                    self.log.insert("end", wert + "\n")
                    self.log.see("end")
                    self.log.configure(state="disabled")
                elif art == "fortschritt":
                    n, t = wert
                    if str(self.fortschritt.cget("mode")) != "determinate":
                        self.fortschritt.stop()
                        self.fortschritt.configure(mode="determinate")
                        self._startzeit = time.time()
                    self.fortschritt.configure(maximum=max(t, 1), value=n)
                    prozent = 100.0 * n / max(t, 1)
                    verstrichen = time.time() - (self._startzeit or time.time())
                    rest = (verstrichen / n * (t - n)) if n else 0
                    self.l_status.config(text=f"{prozent:5.1f} %   –   Datei {n} von {t}")
                    self.l_status2.config(text=f"verstrichen {self._zeit(verstrichen)}   ·   "
                                               f"Restzeit ca. {self._zeit(rest)}")
                    if getattr(self, "popup", None) and self.popup.winfo_exists():
                        self.popup.fortschritt(n, t, verstrichen, rest)
                elif art == "update":
                    self._update_ergebnis(*wert)
                elif art == "update_fertig":
                    self._update_fertig(*wert)
                elif art == "fertig":
                    if getattr(self, "_protokoll", None):
                        self.b_liste.configure(state="normal")
                        self.b_excel.configure(state="normal")
                    self.fortschritt.stop()
                    self.fortschritt.configure(mode="determinate")
                    if wert == "Fertig.":
                        self.fortschritt.configure(value=self.fortschritt.cget("maximum"))
                    self.b_start.configure(state="normal")
                    self.b_stop.configure(state="disabled")
                    gesamt = time.time() - (self._startzeit or time.time())
                    self.l_status.config(text=wert)
                    self.l_status2.config(text=f"Gesamtdauer {self._zeit(gesamt)}" if self._startzeit else "")
                    if getattr(self, "popup", None) and self.popup.winfo_exists():
                        z = self._zaehler or {}
                        kurz = "   ·   ".join(f"{k}: {v}" for k, v in sorted(z.items())[:6])
                        if wert.startswith("Mit Fehler"):
                            kurz = ("Fehler: " + getattr(self, "_letzter_fehler", "unbekannt")[:300]
                                    + "\nDetails: Protokolltext anzeigen (Hauptfenster) oder fehler.log im Programmordner.")
                        self.popup.fertig(wert, gesamt if self._startzeit else 0, kurz)
        except queue.Empty:
            pass
        self.after(150, self._meldungen_abholen)

    @staticmethod
    def _zeit(sek: float) -> str:
        sek = int(max(0, sek))
        h, r = divmod(sek, 3600)
        m, s = divmod(r, 60)
        return f"{h} h {m:02d} min" if h else (f"{m} min {s:02d} s" if m else f"{s} s")

    def _start(self):
        self._einstellungen_uebernehmen()
        e = dict(self.e, _programmordner=str(PROGRAMMORDNER))
        lauf = kern.Sortierlauf(e, log=self._schreiben,
                                fortschritt=lambda n, t: self.meldungen.put(("fortschritt", (n, t))),
                                abbruch=self.abbruch)
        fehler = lauf.pruefen()
        if fehler:
            messagebox.showerror("Eingabe prüfen", fehler)
            return
        if not e["testlaufe"]:
            if not self._alte_ergebnisse_behandeln(Path(e["ziel"]).expanduser(), e.get("doppel_ordner", "")):
                return
        if e["verschieben"] and not e["testlaufe"]:
            if not messagebox.askyesno("Verschieben bestätigen",
                                       "Die Originale werden aus dem Quellordner VERSCHOBEN.\n"
                                       "Das lässt sich nicht automatisch rückgängig machen.\nFortfahren?"):
                return
        self.abbruch.clear()
        self.log.configure(state="normal")
        self.log.delete("1.0", "end")
        self.log.configure(state="disabled")
        self.b_start.configure(state="disabled")
        self.b_stop.configure(state="normal")
        self._startzeit = None
        self.fortschritt.configure(mode="indeterminate", value=0)
        self.fortschritt.start(12)
        vorbereitung = ("Vorbereitung: Dateien suchen" +
                        (", KI-Modell laden (beim ersten Mal Download, kann einige Minuten dauern)"
                         if e["nach_thema"] else "") + " ...")
        self.l_status.config(text=vorbereitung)
        self.l_status2.config(text="")
        self._zaehler = {}
        self.popup = FortschrittFenster(self, self._abbrechen, self._ziel_oeffnen)
        self.popup.vorbereitung(vorbereitung)

        def arbeit():
            try:
                ergebnis = lauf.starten()
                if ergebnis:
                    self._protokoll, self._zaehler = ergebnis[0], ergebnis[1]
                    self._lauf_ziel = Path(e["ziel"]).expanduser()
                    self._lauf_testlauf = bool(e["testlaufe"])
                if getattr(lauf, "netz_abbruch", False):
                    self._letzter_fehler = ("Netzwerkverbindung zum Quell-/Zielordner unterbrochen (NAS, Tailscale). "
                                            "Verbindung prüfen und Lauf erneut starten – bei der Frage nach früheren "
                                            "Ergebnissen NEIN (behalten und ergänzen) wählen.")
                    self.meldungen.put(("fertig", "Mit Fehler beendet."))
                else:
                    self.meldungen.put(("fertig", "Abgebrochen." if self.abbruch.is_set() else "Fertig."))
            except Exception as ex:
                import traceback
                text = traceback.format_exc()
                self._letzter_fehler = f"{type(ex).__name__}: {ex}"
                self._schreiben("\nFEHLER (Lauf abgebrochen):\n" + text)
                try:
                    (PROGRAMMORDNER / "fehler.log").write_text(
                        f"{datetime.now():%Y-%m-%d %H:%M:%S}\n{text}", encoding="utf-8")
                except Exception:
                    pass
                self.meldungen.put(("fertig", "Mit Fehler beendet."))

        self.thread = threading.Thread(target=arbeit, daemon=True)
        self.thread.start()

    def _logtext_umschalten(self):
        if self.v_logtext.get():
            self.log.grid()
            self.f_log.rowconfigure(0, weight=1)
        else:
            self.log.grid_remove()
            self.f_log.rowconfigure(0, weight=0)

    ERGEBNIS_ORDNER = ("nach_Datum", "nach_Ort", "nach_Thema", "Aussortierte_Doppelbilder")

    def _alte_ergebnisse_behandeln(self, ziel: Path, doppel_ordner: str) -> bool:
        """
        Prüft vor einem echten Lauf, ob im Ziel schon Ergebnisse eines früheren Laufs liegen.
        Ja  = alle Ergebnisordner löschen und neu erstellen
        Nein = behalten und ergänzen (bereits vorhandene identische Dateien werden übersprungen)
        Abbrechen = zurück, nichts passiert.   Rückgabe: True = Lauf fortsetzen.
        """
        vorhanden = [ziel / n for n in self.ERGEBNIS_ORDNER if (ziel / n).is_dir()]
        do = Path(doppel_ordner).expanduser() if doppel_ordner.strip() else None
        if do is not None and do.is_dir() and do not in vorhanden and do.name == "Aussortierte_Doppelbilder":
            vorhanden.append(do)
        if not vorhanden:
            return True
        liste = "\n".join(f"   • {p}" for p in vorhanden)
        antwort = messagebox.askyesnocancel(
            "Frühere Ergebnisse gefunden",
            "Im Zielordner liegen bereits Ergebnisse eines früheren Laufs:\n\n" + liste +
            "\n\nWurden seither Themen oder Beschreibungen geändert, sollten die alten Ergebnisse gelöscht "
            "werden, damit keine veralteten Zuordnungen stehen bleiben.\n\n"
            "JA  =  diese Ordner löschen und alles neu erstellen  (die Originale im Quellordner bleiben unberührt)\n"
            "NEIN  =  behalten und nur ergänzen (bereits vorhandene identische Dateien werden übersprungen)\n"
            "ABBRECHEN  =  zurück, nichts wird verändert",
            icon="warning", default="cancel")
        if antwort is None:
            return False
        if antwort:
            if not messagebox.askyesno("Löschen bestätigen",
                                       f"{len(vorhanden)} Ergebnisordner werden jetzt unwiderruflich gelöscht.\n\n"
                                       "Weiter?", icon="warning", default="no"):
                return False
            fehler = []
            for p in vorhanden:
                try:
                    shutil.rmtree(p)
                except Exception as ex:
                    fehler.append(f"{p}: {ex}")
            if fehler:
                messagebox.showerror("Löschen unvollständig",
                                     "Folgende Ordner konnten nicht gelöscht werden (Datei geöffnet? Sync aktiv?):\n\n"
                                     + "\n".join(fehler) + "\n\nDer Lauf wird nicht gestartet.")
                return False
            self._schreiben(f"Alte Ergebnisse gelöscht: {', '.join(p.name for p in vorhanden)}")
        return True

    # ------------------------------------------------------------------ Update
    def _update_pruefen(self, manuell=False):
        if not manuell:
            letzte = float(self.e.get("update_letzte_pruefung", 0) or 0)
            if time.time() - letzte < kern.UPDATE_INTERVALL_H * 3600:
                return
        url = str(self.e.get("update_url", "") or "").strip() or kern.UPDATE_URL
        if "BENUTZER" in url:
            if manuell:
                messagebox.showinfo("Update", "Update-Adresse ist noch nicht eingerichtet (version.json auf GitHub).")
            return
        self.l_update.config(text="Suche nach Updates …")

        def arbeit():
            info = kern.update_pruefen(VERSION, url)
            self.meldungen.put(("update", (info, manuell)))

        threading.Thread(target=arbeit, daemon=True).start()

    def _update_ergebnis(self, info, manuell):
        self.e["update_letzte_pruefung"] = time.time()
        try:
            kern.einstellungen_speichern(EINSTELLUNGSDATEI, self.e)
        except Exception:
            pass
        if not info:
            self.l_update.config(text="Programm ist aktuell." if manuell else "")
            return
        self.l_update.config(text=f"Update {info.get('version')} verfügbar")
        hinweise = str(info.get("hinweise", "")).strip()
        antwort = messagebox.askyesno(
            "Update verfügbar",
            f"Neue Version {info.get('version')} (installiert: {VERSION}).\n\n"
            + (("Änderungen:\n" + hinweise + "\n\n") if hinweise else "")
            + "Jetzt herunterladen und installieren? Das Programm wird danach neu gestartet.\n"
              "Einstellungen und Themen bleiben erhalten.",
            icon="info")
        if not antwort:
            return
        if self.thread and self.thread.is_alive():
            messagebox.showwarning("Update", "Bitte zuerst den laufenden Sortierlauf beenden.")
            return
        self.l_update.config(text="Update wird installiert …")
        self.b_update.configure(state="disabled")

        def arbeit():
            ok, meldung = kern.update_installieren(info, PROGRAMMORDNER, log=lambda t: self.meldungen.put(("log", t)))
            self.meldungen.put(("update_fertig", (ok, meldung)))

        threading.Thread(target=arbeit, daemon=True).start()

    def _update_fertig(self, ok, meldung):
        self.b_update.configure(state="normal")
        if not ok:
            self.l_update.config(text="Update fehlgeschlagen")
            messagebox.showerror("Update fehlgeschlagen", meldung + "\n\nDas bisherige Programm bleibt unverändert.")
            return
        self.l_update.config(text="Update installiert – Neustart …")
        self._einstellungen_uebernehmen()
        start = PROGRAMMORDNER / "Fotos_sortieren_start.pyw"
        try:
            exe = sys.executable
            if sys.platform.startswith("win"):
                pw = Path(exe).with_name("pythonw.exe")
                exe = str(pw) if pw.exists() else exe
            subprocess.Popen([exe, str(start)], cwd=str(PROGRAMMORDNER))
        except Exception as ex:
            messagebox.showinfo("Update installiert", meldung + f"\n\nBitte das Programm von Hand neu starten ({ex}).")
        self.destroy()

    def _liste_zeigen(self):
        if not getattr(self, "_protokoll", None):
            messagebox.showinfo("Hinweis", "Noch kein Lauf ausgeführt.")
            return
        ZuordnungsFenster(self, self._protokoll, self._lauf_ziel, self._lauf_testlauf, PROGRAMMORDNER)

    def _excel_direkt(self):
        if not getattr(self, "_protokoll", None):
            messagebox.showinfo("Hinweis", "Noch kein Lauf ausgeführt.")
            return
        f = ZuordnungsFenster(self, self._protokoll, self._lauf_ziel, self._lauf_testlauf, PROGRAMMORDNER)
        f.excel()

    def _abbrechen(self):
        self.abbruch.set()
        self.l_status.config(text="Wird abgebrochen ...")

    @staticmethod
    def _ordner_im_explorer(pfad: Path):
        """Ordner im Explorer öffnen (Windows: os.startfile, Rückfall explorer.exe)."""
        pfad = Path(pfad)
        try:
            if sys.platform.startswith("win"):
                try:
                    os.startfile(str(pfad))  # type: ignore[attr-defined]
                except Exception:
                    subprocess.Popen(["explorer", str(pfad)])
            elif sys.platform == "darwin":
                subprocess.Popen(["open", str(pfad)])
            else:
                subprocess.Popen(["xdg-open", str(pfad)])
            return True
        except Exception:
            return False

    def _ziel_oeffnen(self):
        """Zielordner öffnen; nach einem Testlauf (Ziel existiert nicht) den Ordner mit Protokoll/Excelliste."""
        eltern = self.popup if getattr(self, "popup", None) and self.popup.winfo_exists() else self
        z = self.v_ziel.get().strip()
        if z and Path(z).is_dir():
            if not self._ordner_im_explorer(Path(z)):
                messagebox.showerror("Fehler", f"Ordner konnte nicht geöffnet werden:\n{z}", parent=eltern)
            return
        # Testlauf: es wurde nichts kopiert -> Protokoll und Excelliste liegen im Programmordner
        if getattr(self, "_lauf_testlauf", False) or not z:
            self._ordner_im_explorer(PROGRAMMORDNER)
            messagebox.showinfo("Testlauf",
                                "Beim Testlauf wird nichts kopiert, der Zielordner existiert deshalb noch nicht.\n\n"
                                f"Geöffnet wurde stattdessen der Programmordner mit Protokoll und Excelliste:\n{PROGRAMMORDNER}",
                                parent=eltern)
        else:
            messagebox.showinfo("Hinweis", f"Zielordner existiert (noch) nicht:\n{z}", parent=eltern)

    def _schliessen(self):
        if self.thread and self.thread.is_alive():
            if not messagebox.askyesno("Beenden", "Ein Lauf ist aktiv. Wirklich beenden?"):
                return
            self.abbruch.set()
        self._einstellungen_uebernehmen()
        self.destroy()


def _fehler_anzeigen(text: str):
    """Fehler sichtbar machen, auch wenn ohne Konsole (pythonw) gestartet wurde."""
    try:
        (PROGRAMMORDNER / "fehler.log").write_text(text, encoding="utf-8")
    except Exception:
        pass
    try:
        wurzel = tk.Tk()
        wurzel.withdraw()
        messagebox.showerror("Fotosortierprogramm – Fehler beim Start",
                             text[-3000:] + "\n\n(auch gespeichert in fehler.log im Programmordner)")
        wurzel.destroy()
    except Exception:
        print(text, file=sys.stderr)


if __name__ == "__main__":
    import traceback
    try:
        App().mainloop()
    except Exception:
        _fehler_anzeigen(traceback.format_exc())
